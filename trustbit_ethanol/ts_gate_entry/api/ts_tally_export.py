# TS Tally Export — .xlsx endpoint (v2.44; Purchase Invoice + single-sheet download v2.76.0)
# Builds ALL FOUR template sheets (Pmt Receipt / JV / Sales / Purchase Invoice) from the
# report's own row-builders and returns one workbook — or just one sheet when
# `only_sheet` is given. Read-only: no DB write of any kind.
import io
import json
import re

import frappe
import openpyxl
from frappe import _
from frappe.utils import getdate
from frappe.utils.xlsxutils import make_xlsx

from trustbit_ethanol.ts_gate_entry.report.ts_tally_export import ts_tally_export as rpt

EXPORT_ROLES = ("Accounts User", "Accounts Manager", "System Manager")
MAX_ROWS = 50000  # per sheet


def _inert(v):
	# openpyxl types any str starting with "=" as a formula that Excel evaluates on
	# open; "+", "-", "@" stay inline strings, so only "=" needs the apostrophe guard.
	return "'" + v if isinstance(v, str) and v.startswith("=") else v


def _cell(v):
	return "" if v is None else _inert(v)


@frappe.whitelist(methods=["POST"])
def export_tally_excel(filters=None, only_sheet=None):
	"""One-click Tally workbook: all 4 sheets in template order (or just `only_sheet`),
	exact header labels, real date cells (System Settings date format), no formulas."""
	frappe.only_for(EXPORT_ROLES)
	if only_sheet and (not isinstance(only_sheet, str) or only_sheet not in rpt.SHEET_ORDER):
		frappe.throw(_("Invalid sheet"))
	names = [only_sheet] if only_sheet else rpt.SHEET_ORDER
	# Builders use raw SQL, so every exported source doctype is gated explicitly (fail-closed:
	# the user must be able to read every sheet the file will contain).
	for name in names:
		frappe.has_permission(rpt.SHEET_DT[name], "read", throw=True)

	if isinstance(filters, str):
		try:
			filters = json.loads(filters)
		except json.JSONDecodeError:
			frappe.throw(_("Invalid filters payload"))
	if not isinstance(filters, dict):
		frappe.throw(_("Filters must be a dict"))

	sheets = rpt.build_all(filters, names)  # validates from/to
	for name, (rows, _notes) in sheets.items():
		if len(rows) > MAX_ROWS:
			frappe.throw(_("Sheet {0} has {1} rows (limit {2}). Please narrow the date range.")
				.format(name, len(rows), MAX_ROWS))

	# make_xlsx PREPENDS each sheet (create_sheet(name, 0)), so iterate in reverse. A normal
	# Workbook must be passed in — the default write-only one cannot be reused.
	wb = openpyxl.Workbook()
	wb.remove(wb.active)
	for name in reversed(names):
		data = [rpt.headers(name)] + [[_cell(c) for c in r] for r in sheets[name][0]]
		make_xlsx(data, name, wb=wb, column_widths=rpt.widths(name))
	# make_xlsx strips HTML AFTER _inert ran, so text like "<b></b>=..." can still land as a
	# live formula -> defuse every formula cell, then save once ourselves.
	for ws in wb.worksheets:
		for row in ws.iter_rows():
			for c in row:
				if c.data_type == "f":
					c.value = "'" + str(c.value)
	xlsx = io.BytesIO()
	wb.save(xlsx)

	company = filters.get("company")
	tag = (frappe.db.get_value("Company", company, "abbr") or company) if company else "ALL"
	tag = re.sub(r"[^A-Za-z0-9_-]+", "_", str(tag))[:20]
	stem = "TS_Tally_" + only_sheet.replace(" ", "_") if only_sheet else "TS_Tally_Export"
	frappe.local.response.filename = "{0}_{1}_{2}_to_{3}.xlsx".format(
		stem, tag, getdate(filters["from_date"]), getdate(filters["to_date"]))
	frappe.local.response.filecontent = xlsx.getvalue()
	frappe.local.response.type = "download"
